from __future__ import annotations

import html
from io import BytesIO
from pathlib import Path
from typing import Any

import pandas as pd
from jinja2 import Template

from insurance_ai.core.models import ReportPayload


HTML_TEMPLATE = Template("""
<!doctype html><html lang="fr"><head><meta charset="utf-8">
<title>{{ title }}</title><style>
body{font-family:Arial,sans-serif;max-width:1100px;margin:40px auto;color:#1f2937;line-height:1.5}
h1,h2{color:#17365d}.kpis{display:flex;flex-wrap:wrap;gap:12px}.kpi{border:1px solid #dbe3ef;border-radius:10px;padding:12px 18px;min-width:180px;background:#f7f9fc}.kpi b{font-size:1.3em}
table{border-collapse:collapse;width:100%;margin:12px 0}th,td{border:1px solid #d6dbe3;padding:7px;font-size:13px}th{background:#2f5597;color:white} .meta{color:#6b7280;font-size:12px}</style></head><body>
<h1>{{ title }}</h1><p>{{ subtitle }}</p><p class="meta">{{ metadata }}</p>
<h2>Synthèse exécutive</h2><p>{{ executive_summary }}</p>
{% if kpis %}<h2>Indicateurs clés</h2><div class="kpis">{% for k in kpis %}<div class="kpi"><span>{{ k.get('label', k.get('code','KPI')) }}</span><br><b>{{ k.get('display', k.get('value','')) }}</b></div>{% endfor %}</div>{% endif %}
{% for t in tables %}<h2>{{ t.title }}</h2>{{ t.html | safe }}{% endfor %}
{% if findings %}<h2>Constats</h2><ul>{% for x in findings %}<li>{{ x }}</li>{% endfor %}</ul>{% endif %}
{% if recommendations %}<h2>Recommandations</h2><ul>{% for x in recommendations %}<li>{{ x }}</li>{% endfor %}</ul>{% endif %}
{% if sources %}<h2>Sources</h2><ul>{% for x in sources %}<li>{{ x }}</li>{% endfor %}</ul>{% endif %}
</body></html>
""")


class ReportGenerator:
    @staticmethod
    def _table_entries(payload: ReportPayload) -> list[dict[str, Any]]:
        entries = []
        for t in payload.tables:
            data = t.get("data")
            if isinstance(data, pd.DataFrame):
                entries.append({"title": t.get("title", "Tableau"), "data": data})
        return entries

    def html(self, payload: ReportPayload) -> bytes:
        tables = [
            {"title": e["title"], "html": e["data"].head(200).to_html(index=False, border=0)}
            for e in self._table_entries(payload)
        ]
        safe_kpis = []
        for k in payload.kpis:
            safe_kpis.append({key: html.escape(str(value)) for key, value in k.items()})
        rendered = HTML_TEMPLATE.render(
            title=html.escape(payload.title),
            subtitle=html.escape(payload.subtitle),
            executive_summary=html.escape(payload.executive_summary),
            kpis=safe_kpis,
            tables=tables,
            findings=[html.escape(x) for x in payload.findings],
            recommendations=[html.escape(x) for x in payload.recommendations],
            sources=[html.escape(x) for x in payload.sources],
            metadata=html.escape(" | ".join(f"{k}: {v}" for k, v in payload.metadata.items())),
        )
        return rendered.encode("utf-8")

    def docx(self, payload: ReportPayload) -> bytes:
        from docx import Document
        from docx.shared import Inches
        doc = Document()
        doc.add_heading(payload.title, 0)
        doc.add_paragraph(payload.subtitle)
        doc.add_heading("Synthèse exécutive", level=1)
        doc.add_paragraph(payload.executive_summary)
        if payload.kpis:
            doc.add_heading("Indicateurs clés", level=1)
            table = doc.add_table(rows=1, cols=3)
            table.style = "Table Grid"
            for i, h in enumerate(["Indicateur", "Valeur", "Formule"]):
                table.rows[0].cells[i].text = h
            for k in payload.kpis:
                cells = table.add_row().cells
                cells[0].text = str(k.get("label", k.get("code", "KPI")))
                cells[1].text = str(k.get("display", k.get("value", "")))
                cells[2].text = str(k.get("formula", ""))
        for entry in self._table_entries(payload):
            doc.add_heading(entry["title"], level=1)
            df = entry["data"].head(100)
            table = doc.add_table(rows=1, cols=len(df.columns))
            table.style = "Table Grid"
            for i, col in enumerate(df.columns):
                table.rows[0].cells[i].text = str(col)
            for _, row in df.iterrows():
                cells = table.add_row().cells
                for i, col in enumerate(df.columns):
                    cells[i].text = str(row[col])
        if payload.findings:
            doc.add_heading("Constats", level=1)
            for item in payload.findings:
                doc.add_paragraph(item, style="List Bullet")
        if payload.recommendations:
            doc.add_heading("Recommandations", level=1)
            for item in payload.recommendations:
                doc.add_paragraph(item, style="List Bullet")
        for chart in payload.charts:
            image = chart.get("png")
            if image:
                doc.add_heading(chart.get("title", "Graphique"), level=1)
                doc.add_picture(BytesIO(image), width=Inches(6.4))
        buf = BytesIO(); doc.save(buf); return buf.getvalue()

    def pptx(self, payload: ReportPayload) -> bytes:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = payload.title
        slide.placeholders[1].text = payload.subtitle
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Synthèse exécutive"
        slide.placeholders[1].text = payload.executive_summary
        if payload.kpis:
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            slide.shapes.title.text = "Indicateurs clés"
            rows, cols = len(payload.kpis) + 1, 3
            table = slide.shapes.add_table(rows, cols, Inches(0.6), Inches(1.5), Inches(8.8), Inches(4.8)).table
            headers = ["Indicateur", "Valeur", "Formule"]
            for j, h in enumerate(headers): table.cell(0,j).text = h
            for i, k in enumerate(payload.kpis, start=1):
                table.cell(i,0).text = str(k.get("label", k.get("code","KPI")))
                table.cell(i,1).text = str(k.get("display", k.get("value","")))
                table.cell(i,2).text = str(k.get("formula", ""))
        for chart in payload.charts[:8]:
            image = chart.get("png")
            if image:
                slide = prs.slides.add_slide(prs.slide_layouts[5])
                slide.shapes.title.text = chart.get("title", "Graphique")
                slide.shapes.add_picture(BytesIO(image), Inches(0.7), Inches(1.5), width=Inches(8.6))
        if payload.findings or payload.recommendations:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = "Constats et recommandations"
            lines = ["CONSTATS"] + [f"• {x}" for x in payload.findings] + ["", "RECOMMANDATIONS"] + [f"• {x}" for x in payload.recommendations]
            slide.placeholders[1].text = "\n".join(lines)
        buf = BytesIO(); prs.save(buf); return buf.getvalue()

    def pdf(self, payload: ReportPayload) -> bytes:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import cm
        from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        buf = BytesIO()
        pdf = SimpleDocTemplate(buf, pagesize=A4, rightMargin=1.5*cm, leftMargin=1.5*cm, topMargin=1.4*cm, bottomMargin=1.4*cm)
        styles = getSampleStyleSheet(); story = []
        story += [Paragraph(html.escape(payload.title), styles["Title"]), Paragraph(html.escape(payload.subtitle), styles["Normal"]), Spacer(1, 10)]
        story += [Paragraph("Synthèse exécutive", styles["Heading1"]), Paragraph(html.escape(payload.executive_summary), styles["BodyText"])]
        if payload.kpis:
            data = [["Indicateur", "Valeur", "Formule"]] + [[str(k.get("label",k.get("code",""))), str(k.get("display",k.get("value",""))), str(k.get("formula",""))] for k in payload.kpis]
            t = Table(data, repeatRows=1, colWidths=[5*cm, 3.5*cm, 8*cm]); t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#2F5597")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("GRID",(0,0),(-1,-1),0.25,colors.grey),("FONTSIZE",(0,0),(-1,-1),8)])); story += [Spacer(1,8), t]
        for chart in payload.charts[:6]:
            if chart.get("png"):
                story += [Spacer(1,12), Paragraph(html.escape(chart.get("title","Graphique")), styles["Heading2"]), Image(BytesIO(chart["png"]), width=17*cm, height=9*cm)]
        if payload.findings:
            story += [Paragraph("Constats", styles["Heading1"])] + [Paragraph("• "+html.escape(x), styles["BodyText"]) for x in payload.findings]
        if payload.recommendations:
            story += [Paragraph("Recommandations", styles["Heading1"])] + [Paragraph("• "+html.escape(x), styles["BodyText"]) for x in payload.recommendations]
        pdf.build(story); return buf.getvalue()


    @staticmethod
    def _safe_excel_df(df: pd.DataFrame) -> pd.DataFrame:
        safe = df.copy()
        for col in safe.columns:
            if safe[col].dtype == object or pd.api.types.is_string_dtype(safe[col]):
                safe[col] = safe[col].map(
                    lambda x: ("'" + x) if isinstance(x, str) and x[:1] in {"=", "+", "-", "@"} else x
                )
        return safe

    def xlsx(self, payload: ReportPayload, curated_df: pd.DataFrame | None = None) -> bytes:
        buf = BytesIO()
        with pd.ExcelWriter(buf, engine="openpyxl") as writer:
            pd.DataFrame(payload.kpis).to_excel(writer, sheet_name="KPI", index=False)
            pd.DataFrame({"Finding": payload.findings}).to_excel(writer, sheet_name="Findings", index=False)
            pd.DataFrame({"Recommendation": payload.recommendations}).to_excel(writer, sheet_name="Recommendations", index=False)
            for idx, entry in enumerate(self._table_entries(payload)[:5], start=1):
                self._safe_excel_df(entry["data"].head(200_000)).to_excel(writer, sheet_name=f"Analysis_{idx}", index=False)
            if curated_df is not None:
                self._safe_excel_df(curated_df.head(1_000_000)).to_excel(writer, sheet_name="Curated_Data", index=False)
        return buf.getvalue()
